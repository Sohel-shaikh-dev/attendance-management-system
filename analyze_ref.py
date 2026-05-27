import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document(r"C:\attendance_web_app real - Copy\reference Black book.docx")

# Analyze styles used
styles_seen = {}
for i, para in enumerate(doc.paragraphs[:300]):
    s = para.style.name
    t = para.text.strip()[:80]
    if not t:
        continue
    if s not in styles_seen:
        styles_seen[s] = []
    if len(styles_seen[s]) < 3:
        styles_seen[s].append(t)

print("=== STYLES USED ===")
for s, examples in styles_seen.items():
    print(f"\nStyle: {s!r}")
    for e in examples:
        print(f"  -> {e}")

# Analyze first 30 non-empty paragraphs in detail
print("\n=== DETAILED PARA ANALYSIS (first 30 non-empty) ===")
count = 0
for para in doc.paragraphs:
    if not para.text.strip():
        continue
    if count >= 30:
        break
    count += 1
    align = str(para.alignment)
    style = para.style.name
    # Get run formatting
    run_info = []
    for run in para.runs:
        if run.text.strip():
            ri = {
                'bold': run.bold,
                'italic': run.italic,
                'size': run.font.size.pt if run.font.size else None,
                'name': run.font.name,
                'color': str(run.font.color.rgb) if run.font.color and run.font.color.type else None,
            }
            run_info.append(ri)
    print(f"\n[{count}] Style={style!r} Align={align}")
    print(f"     Text: {para.text.strip()[:70]!r}")
    if run_info:
        print(f"     Runs: {run_info[:2]}")

# Check page setup
section = doc.sections[0]
print(f"\n=== PAGE SETUP ===")
print(f"Page width:  {section.page_width.inches:.2f} in")
print(f"Page height: {section.page_height.inches:.2f} in")
print(f"Left margin: {section.left_margin.inches:.2f} in")
print(f"Right margin:{section.right_margin.inches:.2f} in")
print(f"Top margin:  {section.top_margin.inches:.2f} in")
print(f"Bottom margin:{section.bottom_margin.inches:.2f} in")
